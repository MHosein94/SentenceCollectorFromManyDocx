from PyQt6.QtWidgets import QWidget, QApplication, QFileDialog, QMessageBox
from PyQt6.QtGui import QIcon, QFontDatabase
from PyQt6.uic import loadUi
from sys import exit, argv
from glob import glob
from docx import Document
import xlsxwriter
from os.path import dirname, join
        
class MainWindow(QWidget):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        
        loadUi(join(baseDir, 'MainWindow.ui'), self)
        self.setWindowIcon(QIcon(join(baseDir, 'Icon.svg')))

        self._moveCenter()
        self._connectSignals()
        self.leSubject.setText('Subject')
    
    # _______________    
    def _moveCenter(self):
        frame = self.frameGeometry()
        centerPoint = self.screen().availableGeometry().center()
        frame.moveCenter(centerPoint)
        self.move(frame.topLeft())
    
    # _______________
    def _connectSignals(self):
        self.btnOpen.clicked.connect(self.openFolder)

    # _______________        
    def openFolder(self):
        try:
            folderName = QFileDialog.getExistingDirectory(self, caption='Choose the containg folder', directory='')
            fileNames = glob(join(folderName, '*.docx'))
            subject = self.leSubject.text()
            numberOfFiles = len(fileNames)
            subjectsList = []
            filesList = []
            
            for fileNumber, fileName in enumerate(fileNames):
                document = Document(fileName)
                for paragraph in document.paragraphs:
                    for run in paragraph.runs:
                        if subject.lower() in run.text.lower():
                            subjectsList.append(run.text.replace(subject, '').replace(': ', ''))
                            # filesList.append(fileName.split('\\')[-1]) # Windows
                            filesList.append(fileName.split('/')[-1]) #Linux
                            
                self.progressBar.setValue(int(((fileNumber+1)/numberOfFiles)*100))

            if len(subjectsList) > 0:
                workbook = xlsxwriter.Workbook(f'{subject}s.xlsx')
                worksheet = workbook.add_worksheet(f'{subject}s')
                headerFormat = workbook.add_format(
                    {
                    'bold': True,
                    'font_size': 14,
                    'fg_color': '#e3f104',
                    'border': 4
                    }
                    )
                worksheet.write(0, 0, subject, headerFormat) # column header
                worksheet.write(0, 1, 'File Name', headerFormat)
                worksheet.set_column(0, 0, 25) # column width
                worksheet.set_column(1, 1, 25)
                for row, sbj in enumerate(subjectsList):
                    worksheet.write(row+1, 0, sbj)
                    worksheet.write(row+1, 1, filesList[row])

                workbook.close()
                QMessageBox.information(self, 'Message', 'Excel File Created!')
                
        except Exception as e:
            print(e)
            
if __name__ == '__main__':
    app = QApplication(argv)
    baseDir = dirname(__file__)
    QFontDatabase.addApplicationFont(join(baseDir, 'Helvetica.ttf'))

    mainWindow = MainWindow()
    mainWindow.show()
        
    exit(app.exec())